from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Relatorio_Tecnico_Ponto_de_Apoio.docx"
LOGO = ROOT / "public" / "brand" / "ponto-de-apoio-logo-green.png"

GREEN = "006B67"
DARK = "12332E"
MUTED = "566B66"
PALE = "E8F3F0"
LIGHT = "F5F8F7"
WHITE = "FFFFFF"
RED = "9B1C1C"
GOLD = "8A6500"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def font_run(run, size=11, bold=False, color=DARK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, *, bold=False, italic=False, color=DARK, size=11, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    font_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        font_run(p.add_run(f"•  {item}"), size=11)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_properties.append(header_marker)
    for i, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], PALE)
        p = table.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font_run(p.add_run(header), size=10, bold=True, color=GREEN)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            font_run(p.add_run(str(value)), size=9.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, title, body, fill=PALE, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    font_run(p.add_run(title), size=11, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    font_run(p2.add_run(body), size=10, color=DARK)


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, before, after, color in (
    ("Heading 1", 16, 18, 10, GREEN),
    ("Heading 2", 13, 14, 7, GREEN),
    ("Heading 3", 12, 10, 5, DARK),
):
    s = styles[style_name]
    s.font.name = "Calibri"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font_run(header.add_run("PONTO DE APOIO  |  RELATÓRIO TÉCNICO"), size=8.5, bold=True, color=MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font_run(footer.add_run("Documento operacional • agosto de 2026"), size=8.5, color=MUTED)

# Capa: editorial_cover com identidade da marca.
add_text(doc, "RELATÓRIO TÉCNICO E OPERACIONAL", bold=True, color=GREEN, size=10, after=24, align=WD_ALIGN_PARAGRAPH.CENTER)
if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    logo_shape = p.add_run().add_picture(str(LOGO), width=Inches(4.1))
    logo_shape._inline.docPr.set("descr", "Logotipo do Ponto de Apoio")
    logo_shape._inline.docPr.set("title", "Ponto de Apoio")
add_text(doc, "Ponto de Apoio", bold=True, color=DARK, size=30, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(doc, "Arquitetura, segurança, implantação e estado atual da plataforma", color=GREEN, size=15, after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(doc, "Catálogo de profissionais verificados com acolhimento inicial assistido por IA", italic=True, color=MUTED, size=11, after=70, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(doc, "Versão 1.2", bold=True, color=DARK, size=11, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(doc, "24 de agosto de 2026", color=MUTED, size=10, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
add_text(doc, "Preparado para a equipe do Ponto de Apoio", color=MUTED, size=9.5, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

page_break(doc)
doc.add_heading("1. Resumo executivo", level=1)
add_text(doc, "O Ponto de Apoio é uma aplicação web preparada para conectar pessoas a psicólogos que passam por um processo administrativo de verificação. A plataforma combina catálogo público, cadastro profissional, autenticação, painel administrativo e um canal de acolhimento inicial com inteligência artificial.")
add_callout(doc, "Estado atual", "A plataforma está publicada e operacional no domínio pontodeapoio.social.br. Supabase/PostgreSQL com RLS, autenticação, cadastro profissional, painel administrativo, catálogo filtrado, busca de profissionais pelo chat, DNS, Resend e SMTP de recuperação de senha foram configurados e validados em produção.")
doc.add_heading("Escopo entregue", level=2)
add_bullets(doc, [
    "Landing page e identidade visual oficial do Ponto de Apoio.",
    "Chat de acolhimento integrado à OpenAI por rota segura no servidor, sem persistência de conversas.",
    "Busca assistida que consulta somente profissionais approved e publicados, com filtros de modalidade, cidade e UF.",
    "Cadastro e autenticação de psicólogos por e-mail e senha no Supabase Auth.",
    "Área profissional para envio e acompanhamento do cadastro.",
    "Região do CRP selecionada entre as 24 regiões oficiais e UF escolhida entre as 27 unidades federativas.",
    "Painel administrativo para aprovar, solicitar revisão, suspender, restaurar e excluir cadastros.",
    "Catálogo público alimentado somente por profissionais aprovados e publicados.",
    "PostgreSQL relacional, Row Level Security e trilha imutável de auditoria administrativa.",
])
doc.add_heading("Limites importantes", level=2)
add_text(doc, "A IA não realiza diagnóstico e não substitui atendimento profissional. Ela pode apresentar opções compatíveis do catálogo aprovado, mas não realiza recomendação clínica, ranking, agendamento nem confirmação de disponibilidade. O produto ainda não implementa prontuário, agenda, pagamentos, validação automática do registro junto ao CFP/CRP, upload de documentos ou armazenamento das conversas de acolhimento.")

page_break(doc)
doc.add_heading("2. Arquitetura da solução", level=1)
add_table(doc, ["Camada", "Responsabilidade", "Tecnologia"], [
    ("Interface", "Páginas públicas e áreas protegidas", "Next.js App Router + React"),
    ("Servidor web", "Server Actions, APIs, validação de sessão e segredos", "Next.js / Node.js 22"),
    ("Identidade", "Cadastro, login, recuperação e sessões", "Supabase Auth + @supabase/ssr"),
    ("Dados", "Perfis, profissionais e auditoria", "Supabase PostgreSQL"),
    ("Autorização", "Acesso por usuário, papel e estado", "PostgreSQL RLS + grants"),
    ("IA", "Acolhimento inicial e busca controlada no catálogo", "OpenAI Responses API no servidor"),
    ("Hospedagem", "Build, deploy e domínio", "Vercel"),
    ("E-mail", "Confirmação e recuperação de senha", "Resend via SMTP do Supabase"),
], [1800, 4200, 3360])
doc.add_heading("Fluxo lógico", level=2)
add_callout(doc, "Navegador → Next.js → Supabase", "O navegador usa somente a URL e a chave publicável do Supabase. Operações protegidas validam a sessão no servidor. O PostgreSQL aplica privilégios e RLS mesmo quando a requisição passa pela API do Supabase.")
add_callout(doc, "Navegador → /api/chat → OpenAI", "A chave OPENAI_API_KEY permanece no servidor. O frontend envia mensagens à rota interna; a aplicação não grava as conversas no banco nesta versão.", fill=LIGHT)
add_callout(doc, "OpenAI → buscar_profissionais → Supabase", "Quando a pessoa pede opções de atendimento, a OpenAI pode solicitar uma ferramenta interna. O servidor valida modalidade, cidade e UF e consulta somente status=approved com is_published=true. A RLS reafirma essa restrição no PostgreSQL.", fill=LIGHT)
page_break(doc)
doc.add_heading("Organização do código", level=2)
add_table(doc, ["Diretório", "Conteúdo"], [
    ("src/app", "Rotas do App Router, Server Actions e APIs"),
    ("src/components", "Componentes de interface por domínio"),
    ("src/lib/supabase", "Clientes browser, servidor e proxy de sessão"),
    ("src/types", "Tipos TypeScript do banco e da aplicação"),
    ("supabase/migrations", "Evolução versionada do PostgreSQL"),
    ("public/brand", "Logotipos e símbolo oficial"),
    ("docs", "Documentação técnica e operacional"),
], [2700, 6660])

page_break(doc)
doc.add_heading("3. Tecnologias e versões", level=1)
add_table(doc, ["Tecnologia", "Versão/linha", "Uso"], [
    ("Next.js", "16.3.x", "Framework web e App Router"),
    ("React", "19.2.x", "Interface e componentes"),
    ("TypeScript", "5.9.x", "Tipagem estática"),
    ("Node.js", ">= 22", "Runtime de build e servidor"),
    ("Supabase JS", "2.112.x", "API de Auth e dados"),
    ("Supabase SSR", "0.7.x", "Sessões por cookies no Next.js"),
    ("OpenAI SDK", "7.5.x", "Integração do acolhimento"),
    ("Tailwind CSS", "4.3.x", "Infraestrutura de estilos"),
    ("ESLint / Prettier", "9.x / 3.9.x", "Qualidade e formatação"),
    ("PostgreSQL", "Gerenciado pelo Supabase", "Banco relacional e políticas RLS"),
], [2200, 1900, 5260])
doc.add_heading("Variáveis de ambiente", level=2)
add_table(doc, ["Variável", "Exposição", "Finalidade"], [
    ("NEXT_PUBLIC_SUPABASE_URL", "Pública", "Endpoint do projeto Supabase"),
    ("NEXT_PUBLIC_SUPABASE_PUBLISHABLE_KEY", "Pública", "Chave publicável sujeita a RLS"),
    ("NEXT_PUBLIC_SITE_URL", "Pública", "URL canônica para redirecionamentos"),
    ("OPENAI_API_KEY", "Segredo", "Chave usada somente em /api/chat"),
    ("Senha SMTP / API key Resend", "Segredo", "Configurada no Supabase, nunca no navegador"),
], [3400, 1400, 4560])
add_callout(doc, "Regra de segurança", "Nunca versionar service_role, chave Secret do Supabase, senha do PostgreSQL, chave da OpenAI, senha SMTP ou tokens pessoais. A chave publicável não substitui RLS; ela apenas identifica o projeto.", fill="FFF5E6", accent=GOLD)

page_break(doc)
doc.add_heading("4. Banco de dados e segurança", level=1)
doc.add_heading("Modelo relacional", level=2)
add_table(doc, ["Tabela", "Responsabilidade", "Proteção principal"], [
    ("profiles", "Nome e papel ligados ao auth.users", "Usuário vê o próprio perfil; role não é controlável pelo cliente"),
    ("professionals", "CRP, modalidade, localização, contato e apresentação", "Proprietário edita estados permitidos; admin decide publicação"),
    ("admin_audit_logs", "Registro das decisões administrativas", "Leitura/criação administrativa; sem alteração ou exclusão pela aplicação"),
], [2000, 3500, 3860])
doc.add_heading("Estados do cadastro", level=2)
add_text(doc, "draft → pending_review → approved → suspended. Um cadastro rejeitado pode ser corrigido e reenviado. Somente approved com is_published=true aparece no catálogo.")
doc.add_heading("Controles implementados", level=2)
add_bullets(doc, [
    "RLS habilitado desde a migration inicial.",
    "Grants explícitos para os papéis anon e authenticated.",
    "Função is_admin com security definer e search_path vazio.",
    "Administrador não pode ser criado pela interface pública.",
    "Ações administrativas registradas em admin_audit_logs.",
    "Health check consulta o banco sem retornar linhas ou detalhes internos.",
    "Dados profissionais separados de qualquer futuro dado de acolhimento.",
    "OpenAI sem acesso direto ao banco; a ferramenta recebe no máximo oito resultados com campos públicos mínimos.",
])
add_callout(doc, "Privacidade e dados enviados à IA", "Não existe tabela de conversas e o conteúdo do chat não é persistido. A ferramenta envia somente nome público, CRP, apresentação, modalidade, cidade, UF e URL do catálogo. E-mail, UUID, autenticação, estado administrativo, auditoria e conteúdo de acolhimento não são incluídos. Uma futura persistência exigirá finalidade, consentimento, retenção e revisão jurídica/LGPD.")

page_break(doc)
doc.add_heading("5. Jornadas da aplicação", level=1)
doc.add_heading("Visitante", level=2)
add_bullets(doc, ["Consulta a apresentação institucional.", "Visualiza somente profissionais aprovados e publicados.", "Pode usar o acolhimento inicial, com avisos de segurança e emergência.", "Pode pedir opções por modalidade, cidade e UF; recebe resultados em ordem neutra, sem ranking clínico."])
doc.add_heading("Psicólogo", level=2)
add_bullets(doc, ["Cria conta em /cadastro-profissional e confirma o e-mail.", "Entra em /entrar e acessa /area-profissional.", "Preenche número do CRP, escolhe uma das 24 regiões oficiais, modalidade, cidade, UF, contato e apresentação.", "Envia o cadastro para análise; não consegue autopublicar o perfil."])
doc.add_heading("Equipe administrativa", level=2)
add_bullets(doc, ["Entra pela mesma rota, mas o papel admin direciona a /admin.", "Analisa todos os cadastros permitidos pelas políticas.", "Pode aprovar, pedir revisão, suspender, restaurar ou excluir o registro profissional.", "Deve preferir suspensão para impedimento temporário, preservando o histórico."])
doc.add_heading("Recuperação de senha", level=2)
add_text(doc, "A aplicação possui /recuperar-senha, /auth/callback e /definir-senha. O callback aceita token_hash do tipo recovery e restringe o parâmetro next a caminhos internos, o que permite abrir o e-mail em outro navegador sem criar redirecionamento aberto.")
add_callout(doc, "Configuração validada em produção", "O template Reset Password do Supabase aponta para /auth/callback com token_hash. O subdomínio auth.pontodeapoio.social.br está verificado no Resend, o SMTP próprio está ativo no Supabase e o envio e a abertura de um e-mail real de recuperação foram testados.")

page_break(doc)
doc.add_heading("6. Infraestrutura, domínio e publicação", level=1)
add_table(doc, ["Componente", "Configuração"], [
    ("Código-fonte", "GitHub: grupoentra21/ponto-de-apoio"),
    ("Branch de produção", "main"),
    ("Hospedagem", "Projeto ponto-de-apoio na Vercel"),
    ("URL de contingência", "https://ponto-de-apoio.vercel.app"),
    ("Domínio próprio", "pontodeapoio.social.br"),
    ("Subdomínio de e-mail", "auth.pontodeapoio.social.br"),
    ("Banco/Auth", "Projeto Supabase ponto-de-apoio"),
], [2700, 6660])
doc.add_heading("DNS e e-mail transacional", level=2)
add_table(doc, ["Tipo", "Nome", "Finalidade"], [
    ("A", "@", "Apontar o domínio principal para a Vercel"),
    ("TXT", "_vercel", "Verificação de posse na Vercel"),
    ("TXT", "resend._domainkey.auth", "DKIM do Resend"),
    ("CNAME", "rsend.auth", "Retorno e autenticação do envio"),
    ("CNAME", "send.auth", "Roteamento de envio do Resend"),
], [1100, 3000, 5260])
add_callout(doc, "Situação em 24/08/2026", "Os registros DNS foram inseridos e propagados. O domínio principal está validado na Vercel e serve a produção; o subdomínio de autenticação está validado no Resend; a credencial SMTP está configurada no Supabase; e o fluxo real de recuperação de senha foi concluído com sucesso.")

page_break(doc)
doc.add_heading("7. Histórico das principais atualizações", level=1)
add_table(doc, ["Referência", "Entrega"], [
    ("PR #1", "Fundação Next.js e Supabase"),
    ("90a7d16", "Clientes Supabase e health check seguro"),
    ("PR #2", "Chat de acolhimento integrado à OpenAI"),
    ("PR #3", "Área administrativa e cadastro de psicólogos"),
    ("PR #4", "Identidade visual oficial e assets de marca"),
    ("PR #5", "Fluxo seguro de recuperação de senha"),
    ("2049fe4", "Callback token_hash e grants documentados do banco"),
    ("458f432", "Correção da validação do formato da região do CRP"),
    ("24fb9a8", "Listas oficiais de regiões do CRP e unidades federativas"),
    ("f9e2648", "Busca segura de profissionais aprovados pelo chat"),
    ("a0f9c41", "Documentação inicial da busca assistida"),
], [1900, 7460])
doc.add_heading("Rotas publicadas", level=2)
add_table(doc, ["Rota", "Função", "Acesso"], [
    ("/", "Página institucional", "Público"),
    ("/acolhimento", "Acolhimento assistido por IA", "Público"),
    ("/profissionais", "Catálogo verificado", "Público"),
    ("/cadastro-profissional", "Criação de conta", "Público"),
    ("/entrar", "Login", "Público"),
    ("/area-profissional", "Cadastro e status", "Autenticado"),
    ("/admin", "Operação do catálogo", "Administrador"),
    ("/api/health/supabase", "Verificação do banco", "Servidor"),
    ("/api/chat", "OpenAI e busca controlada no catálogo", "Servidor"),
], [2600, 4300, 2460])

page_break(doc)
doc.add_heading("8. Verificação, operação e próximos passos", level=1)
doc.add_heading("Qualidade executada", level=2)
add_bullets(doc, [
    "npm ci: dependências instaladas com lockfile reproduzível e sem vulnerabilidades reportadas.",
    "Prettier: arquivos alterados nesta entrega aprovados.",
    "npm run lint: ESLint aprovado.",
    "npm run build: compilação e TypeScript aprovados.",
    "Health check de produção: resposta HTTP 200 com status ok.",
])
doc.add_heading("Validações de produção concluídas", level=2)
for number, item in enumerate([
    "Domínio pontodeapoio.social.br apontado para a Vercel e marcado como produção.",
    "Subdomínio auth.pontodeapoio.social.br verificado no Resend.",
    "SMTP do Resend configurado no Supabase sem exposição da credencial no navegador.",
    "Site URL, redirect URLs e template de recuperação atualizados no Supabase.",
    "E-mail real de recuperação entregue, aberto e validado no domínio próprio.",
    "Cadastro profissional testado com status pending_review e oculto do catálogo público.",
    "Listas de região do CRP e UF publicadas e verificadas no deploy de produção.",
    "Busca pelo chat testada em produção com retorno exclusivo de profissional aprovado e publicado, sem dados privados.",
], start=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font_run(p.add_run(f"{number}.  {item}"), size=11)
page_break(doc)
doc.add_heading("Recomendações antes da operação comercial", level=2)
add_bullets(doc, [
    "Concluir a migration e o deploy do recurso de foto profissional otimizada antes de anunciá-lo como disponível.",
    "Revisão jurídica, LGPD e regras do Conselho Federal/Regional de Psicologia.",
    "MFA obrigatório para administradores e revisão periódica de acessos.",
    "Processo documentado de validação e revalidação do CRP.",
    "Testes automatizados de RLS, autenticação, Server Actions e fluxos críticos.",
    "Testes automatizados da ferramenta de busca, incluindo estados inelegíveis e respostas vazias.",
    "Monitoramento de qualidade, equidade e possíveis vieses na apresentação de profissionais.",
    "Monitoramento, alertas, resposta a incidentes e estratégia de backup.",
    "Modelagem separada para pagamentos, agenda e qualquer dado sensível futuro.",
])
add_callout(doc, "Conclusão", "A fundação técnica, o domínio próprio, o e-mail transacional e a busca assistida no catálogo estão operacionais. A indicação é apresentada como compatibilidade objetiva, sem diagnóstico ou ranking clínico. O próximo ciclo deve priorizar foto profissional segura, validação operacional de CRP, testes automatizados, governança LGPD e controles comerciais.")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
